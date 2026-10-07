import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_quotation_to_excel(quotation, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    
    # Headers
    headers = ["S.N.", "Particulars", "Size (Sq.ft)", "Total Size", "Rate (\u20b9)", "Amount (\u20b9)"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num+1)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        
    current_row = 5
    section_index = 0
    section_labels = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    
    for section in quotation.sections.all():
        # Section Header
        ws.cell(row=current_row, column=2, value=section_labels[section_index])
        ws.cell(row=current_row, column=3, value=section.name)
        
        # Merge columns D to G for section header
        ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=7)
        current_row += 1
        section_index += 1
        
        item_counter = 1
        section_start_row = current_row
        
        for item in section.items.all():
            ws.cell(row=current_row, column=2, value=item_counter)
            ws.cell(row=current_row, column=3, value=item.description)
            
            if item.calculation_type == 'AREA_RATE':
                size_str = f"{item.length}*{item.width}" if item.length and item.width else ""
                ws.cell(row=current_row, column=4, value=size_str)
                # Re-inject the native Excel formula
                ws.cell(row=current_row, column=5, value=f'=VALUE(LEFT(D{current_row}, FIND("*", D{current_row}) - 1)) * VALUE(MID(D{current_row}, FIND("*", D{current_row}) + 1, LEN(D{current_row})))')
            else:
                ws.cell(row=current_row, column=4, value="")
                ws.cell(row=current_row, column=5, value=float(item.quantity) if item.quantity else 1.0)
                
            ws.cell(row=current_row, column=6, value=float(item.rate))
            # Amount formula
            ws.cell(row=current_row, column=7, value=f'=(F{current_row}*E{current_row})')
            
            current_row += 1
            item_counter += 1
            
        # Empty row after section
        current_row += 1

    # Grand Total
    ws.cell(row=current_row, column=4, value="Total")
    ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=6)
    # The total sum formula could be injected here, but for simplicity we can just write the backend calculated value
    # or create a complex sum formula based on the section rows
    ws.cell(row=current_row, column=7, value=float(quotation.total_amount))
    ws.cell(row=current_row, column=7).font = Font(bold=True)

    wb.save(filepath)
    return filepath
