import openpyxl
from decimal import Decimal
from quotations.models import Quotation, QuotationSection, QuotationItem
from quotations.services.calculation import calculate_item, calculate_section, calculate_quotation

def import_excel_quotation(filepath, project, quotation_number, date):
    """
    Imports the legacy Quotation format from Excel.
    Reads D:\MY\Dad Work\Common Temp.xlsx structure.
    """
    wb = openpyxl.load_workbook(filepath, data_only=False)
    sheet = wb.active # Usually Sheet1
    
    q = Quotation.objects.create(
        project=project,
        quotation_number=quotation_number,
        date=date
    )
    
    current_section = None
    item_order = 0
    section_order = 0
    
    for row_idx in range(5, sheet.max_row):
        sn = sheet.cell(row=row_idx, column=2).value
        particulars = sheet.cell(row=row_idx, column=3).value
        
        # If S.N. is A, B, C, D... it's a section header
        if sn and isinstance(sn, str) and sn.isalpha() and len(sn) == 1:
            if particulars and particulars.lower() != 'total':
                current_section = QuotationSection.objects.create(
                    quotation=q,
                    name=str(particulars).strip(),
                    order=section_order
                )
                section_order += 1
                item_order = 0
                continue
                
        # If S.N. is a number, it's an item
        if current_section and sn and isinstance(sn, int):
            size_str = sheet.cell(row=row_idx, column=4).value
            rate_val = sheet.cell(row=row_idx, column=6).value
            
            item = QuotationItem(
                section=current_section,
                description=str(particulars).strip() if particulars else "Item",
                order=item_order
            )
            item_order += 1
            
            rate_dec = Decimal('0')
            if rate_val is not None:
                try:
                    rate_dec = Decimal(str(rate_val))
                except:
                    pass
            item.rate = rate_dec
            
            # Parse size_str which looks like '8.5*8.3'
            if size_str and isinstance(size_str, str) and '*' in size_str:
                parts = size_str.split('*')
                if len(parts) == 2:
                    try:
                        item.length = Decimal(parts[0].strip())
                        item.width = Decimal(parts[1].strip())
                        item.calculation_type = 'AREA_RATE'
                    except:
                        item.calculation_type = 'FIXED'
            elif size_str and (isinstance(size_str, int) or isinstance(size_str, float)):
                item.quantity = Decimal(str(size_str))
                item.calculation_type = 'QTY_RATE'
            else:
                item.calculation_type = 'FIXED'
                
            calculate_item(item)
            item.save()
            
    # Recalculate totals
    for sec in q.sections.all():
        calculate_section(sec)
        sec.save()
        
    calculate_quotation(q)
    return q
