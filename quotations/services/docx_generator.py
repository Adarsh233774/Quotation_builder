from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx_quotation(quotation, filepath):
    doc = Document()
    
    # Title
    title = doc.add_heading('QUOTATION', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Meta Info
    doc.add_paragraph(f"Quotation No: {quotation.quotation_number}")
    doc.add_paragraph(f"Date: {quotation.date}")
    doc.add_paragraph(f"Project: {quotation.project.name}")
    doc.add_paragraph(f"Client: {quotation.project.client.name}")
    doc.add_paragraph("")
    
    # Sections
    for section in quotation.sections.all():
        sec_head = doc.add_heading(section.name, level=1)
        
        # Table
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Shading Accent 1'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'S.N.'
        hdr_cells[1].text = 'Description'
        hdr_cells[2].text = 'Qty / Size'
        hdr_cells[3].text = 'Rate'
        hdr_cells[4].text = 'Amount'
        
        for idx, item in enumerate(section.items.all(), 1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(idx)
            row_cells[1].text = item.description or ''
            
            size_str = ""
            if item.calculation_type == 'AREA_RATE':
                size_str = f"{item.length} x {item.width}"
            else:
                size_str = str(item.quantity)
                
            row_cells[2].text = size_str
            row_cells[3].text = str(item.rate)
            row_cells[4].text = str(item.amount)
            
        doc.add_paragraph(f"Section Total: Rs. {section.section_total}")
        doc.add_paragraph("")
        
    doc.add_heading('INVESTMENT SUMMARY', level=1)
    doc.add_paragraph(f"Subtotal: Rs. {quotation.subtotal}")
    if quotation.tax:
        doc.add_paragraph(f"Tax: Rs. {quotation.tax}")
    if quotation.additional_charges:
        doc.add_paragraph(f"Additional Charges: Rs. {quotation.additional_charges}")
    
    p = doc.add_paragraph(f"Total Investment: Rs. {quotation.total_amount}")
    p.runs[0].bold = True
    
    doc.save(filepath)
    return filepath
